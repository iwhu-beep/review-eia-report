"""
从 Word (.docx) / PDF 中提取文本内容，用于环评报告表审核。

用法：
    python scripts/extract_docx.py "<文件路径>" [--format md|json|txt] [--out 结果文件] [--max-chars N] [--include-tables/--no-tables]

向后兼容：python scripts/extract_docx.py file.docx 仍可工作（默认 md 输出到 stdout）。

依赖：pip install -r scripts/requirements.txt
"""

import argparse
import datetime
import json
import os
import sys


def eprint(*args, **kwargs):
    print(*args, file=sys.stderr, **kwargs)


def load_docx(path):
    try:
        from docx import Document
    except ImportError:
        eprint("错误：缺少 python-docx，请执行 pip install -r scripts/requirements.txt")
        sys.exit(2)
    return Document(path)


def extract_docx_meta(doc, path):
    props = doc.core_properties
    meta = {
        "file": os.path.basename(path),
        "title": getattr(props, "title", "") or "",
        "author": getattr(props, "author", "") or "",
        "created": str(getattr(props, "created", "") or ""),
        "modified": str(getattr(props, "modified", "") or ""),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "extracted_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    return meta


def is_heading(paragraph):
    try:
        style_name = paragraph.style.name if paragraph.style else ""
    except Exception:
        style_name = ""
    if "Heading" in style_name or "标题" in style_name:
        return True
    text = paragraph.text.strip()
    # 兜底：中文章节号开头且短句视为标题
    if len(text) < 40 and text[:2].strip() and any(
        text.startswith(p) for p in ("一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十、", "1.", "2.", "3.")
    ):
        return True
    return False


def extract_docx(path, include_tables=True):
    doc = load_docx(path)
    meta = extract_docx_meta(doc, path)

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        paragraphs.append({"text": text, "is_heading": is_heading(para)})

    headers_footers = []
    try:
        for section in doc.sections:
            for part in (section.header, section.footer):
                for para in part.paragraphs:
                    t = para.text.strip()
                    if t and t not in headers_footers:
                        headers_footers.append(t)
    except Exception as exc:
        eprint(f"提示：页眉页脚提取跳过（{exc}）")

    tables = []
    if include_tables:
        for t_idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    c = cell.text.strip().replace("\n", " / ").replace("\r", "")
                    cells.append(c)
                rows.append(cells)
            tables.append({"index": t_idx + 1, "rows": rows, "n_rows": len(rows)})

    comments = []
    try:
        # python-docx 对批注支持有限，尝试读取 comments part，失败则忽略
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        for rel in doc.part.rels.values():
            if "comments" in (rel.target_ref or ""):
                comments.append(f"文档含批注部件：{rel.target_ref}（请在 Word 中查看批注）")
                break
    except Exception:
        pass

    return {"meta": meta, "headers_footers": headers_footers, "paragraphs": paragraphs, "tables": tables, "comments": comments}


def extract_pdf(path, include_tables=True):
    try:
        from pypdf import PdfReader
    except ImportError:
        eprint("错误：PDF 解析需要 pypdf，请执行 pip install pypdf 后重试；扫描件 PDF 请先 OCR。")
        sys.exit(2)
    reader = PdfReader(path)
    paragraphs = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            eprint(f"提示：第 {i + 1} 页解析失败，已跳过（{exc}）")
            continue
        for line in text.splitlines():
            line = line.strip()
            if line:
                paragraphs.append({"text": line, "is_heading": False, "page": i + 1})
    meta = {
        "file": os.path.basename(path),
        "pages": len(reader.pages),
        "paragraph_count": len(paragraphs),
        "table_count": 0,
        "note": "PDF 表格按文本行提取，未做结构化还原" if include_tables else "",
        "extracted_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    return {"meta": meta, "headers_footers": [], "paragraphs": paragraphs, "tables": [], "comments": []}


def to_markdown(data):
    out = []
    meta = data.get("meta", {})
    out.append(f"# 文档提取结果：{meta.get('file', '')}\n")
    out.append("## 元数据")
    for k, v in meta.items():
        out.append(f"- {k}: {v}")
    if data.get("headers_footers"):
        out.append("\n## 页眉页脚")
        for h in data["headers_footers"]:
            out.append(f"- {h}")
    out.append("\n=== 正文内容 ===\n")
    for p in data.get("paragraphs", []):
        t = p["text"]
        out.append(f"\n【{t}】" if p.get("is_heading") else t)
    if data.get("tables"):
        out.append("\n\n=== 表格内容 ===\n")
        for tb in data["tables"]:
            out.append(f"\n--- 表格 {tb['index']}（{tb['n_rows']}行） ---")
            for row in tb["rows"]:
                out.append(" | ".join(row))
    if data.get("comments"):
        out.append("\n\n=== 批注提示 ===\n")
        out.extend(f"- {c}" for c in data["comments"])
    return "\n".join(out)


def to_text(data):
    # 精简纯文本：标题+正文+表格行，适合直接粘贴给模型
    lines = []
    for p in data.get("paragraphs", []):
        lines.append(p["text"])
    for tb in data.get("tables", []):
        lines.append(f"[表格{tb['index']}]")
        for row in tb["rows"]:
            lines.append(" | ".join(row))
    return "\n".join(lines)


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="提取环评报告表文档文本（docx/pdf）")
    ap.add_argument("input", help="输入文件路径（.docx/.pdf）")
    ap.add_argument("--format", choices=["md", "json", "txt"], default="md", help="输出格式（默认md）")
    ap.add_argument("--out", default="", help="输出文件路径，缺省输出到stdout")
    ap.add_argument("--max-chars", type=int, default=0, help="截断提示：超过N字符时在末尾追加分块提示（0=不截断）")
    ap.add_argument("--include-tables", dest="include_tables", action="store_true", default=True)
    ap.add_argument("--no-tables", dest="include_tables", action="store_false", help="不提取表格")
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    path = args.input

    if not os.path.exists(path):
        eprint(f"错误：文件不存在 - {path}")
        sys.exit(2)

    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        data = extract_docx(path, include_tables=args.include_tables)
    elif ext == ".pdf":
        data = extract_pdf(path, include_tables=args.include_tables)
    elif ext == ".doc":
        eprint("错误：.doc 为老格式，请先用 Word/WPS 另存为 .docx 后再提取。")
        sys.exit(2)
    else:
        eprint(f"错误：不支持的文件类型 {ext}，仅支持 .docx/.pdf。")
        sys.exit(2)

    if not data.get("paragraphs") and not data.get("tables"):
        eprint("错误：文档为空或无法提取到文本（扫描件PDF请先OCR）。")
        sys.exit(2)

    if args.format == "json":
        output = json.dumps(data, ensure_ascii=False, indent=2)
    elif args.format == "txt":
        output = to_text(data)
    else:
        output = to_markdown(data)

    if args.max_chars and len(output) > args.max_chars:
        output = output[: args.max_chars]
        output += (
            f"\n\n…[截断：全文超 {args.max_chars} 字符，已截断。"
            "建议按 SKILL.md 长文档策略分章节提取审核，表格块单独完整提取。]"
        )

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(output)
        print(f"已写入：{args.out}（{len(output)}字符）")
    else:
        print(output)


if __name__ == "__main__":
    main()
